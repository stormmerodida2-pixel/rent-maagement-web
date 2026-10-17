"""Helpers for payments: encrypted model fields and M-Pesa statement uploads. (The live M-Pesa connection
is in mpesa.py.)"""

import base64
import csv
import datetime
import hashlib
import io
import logging
import os
import re
from decimal import Decimal, InvalidOperation

from cryptography.fernet import Fernet, InvalidToken
from django.conf import settings
from django.core.files.storage import FileSystemStorage
from django.db import models, transaction
from django.utils import timezone


# --- Private files --------------------------------------------------------------------------------------

class PrivateStorage(FileSystemStorage):
    """Receipts and meter photos: outside MEDIA_ROOT, so the web server never serves them directly. The folder is
    read from settings each time, so it follows PRIVATE_MEDIA_ROOT (tests point it at a temporary folder)."""

    @property
    def base_location(self):
        return settings.PRIVATE_MEDIA_ROOT

    @property
    def location(self):
        return os.path.abspath(self.base_location)

    @property
    def base_url(self):
        return None


def private_storage():
    return PrivateStorage()


# --- Encrypted model fields --------------------------------------------------------------------------

logger = logging.getLogger(__name__)


def _fernet():
    key = settings.FIELD_ENCRYPTION_KEY
    if not key:
        # Fall back to a key derived from SECRET_KEY. Set FIELD_ENCRYPTION_KEY in production
        # so rotating SECRET_KEY doesn't make stored credentials unreadable.
        key = base64.urlsafe_b64encode(hashlib.sha256(settings.SECRET_KEY.encode()).digest())
    return Fernet(key)


class EncryptedTextField(models.TextField):
    """Text encrypted at rest with Fernet. Not usable in lookups."""

    def from_db_value(self, value, expression, connection):
        if not value:
            return value
        try:
            return _fernet().decrypt(value.encode()).decode()
        except InvalidToken:
            logger.error('Could not decrypt %s; was the encryption key changed?', self.name)
            return ''

    def get_prep_value(self, value):
        value = super().get_prep_value(value)
        if not value:
            return value
        return _fernet().encrypt(value.encode()).decode()


# --- M-Pesa statement uploads -------------------------------------------------------------------------
# Landlords without a Daraja app (or catching up on missed payments) download their Paybill/Till statement
# from Safaricom's M-Pesa business portal and upload it. Money paid in is matched to tenants the same way
# live payments are: by account number, then by phone.

STATEMENT_MAX_BYTES = 5 * 1024 * 1024
STATEMENT_MAX_ROWS = 5000
STATEMENT_COLUMNS = {
    'receipt': ('receipt no.', 'receipt no', 'receipt', 'receipt number', 'transaction id', 'trans id'),
    'time': ('completion time', 'transaction time', 'completion date', 'date', 'time'),
    'details': ('details', 'description', 'narration'),
    'status': ('transaction status', 'status'),
    'paid_in': ('paid in', 'credit', 'amount', 'paid in (ksh)'),
    'party': ('other party info', 'other party', 'customer', 'payer'),
    'account': ('a/c no.', 'a/c no', 'account no.', 'account no', 'account number', 'bill ref', 'bill ref no'),
}
TIME_FORMATS = ('%Y-%m-%d %H:%M:%S', '%d-%m-%Y %H:%M:%S', '%d/%m/%Y %H:%M:%S', '%Y/%m/%d %H:%M:%S',
                '%Y-%m-%d %H:%M', '%d-%m-%Y %H:%M', '%d/%m/%Y %H:%M', '%d.%m.%Y %H:%M:%S')


class StatementError(Exception):
    """A problem with the uploaded file as a whole, in words the landlord can act on."""


def _table(upload):
    name = (upload.name or '').lower()
    if name.endswith('.xls'):
        raise StatementError('That is an old-style Excel file (.xls). Open it in Excel and use Save As → '
                             '"Excel Workbook (.xlsx)", or download the statement as CSV, then upload that.')
    if name.endswith('.pdf'):
        raise StatementError("PDF statements can't be read. On the M-Pesa business portal, download the "
                             'statement as Excel or CSV instead.')
    if name.endswith('.csv'):
        return list(csv.reader(io.StringIO(upload.read().decode('utf-8-sig', errors='replace'))))
    if name.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(upload, read_only=True, data_only=True).worksheets[0]
        except Exception as exc:
            raise StatementError("That file couldn't be opened as an Excel workbook.") from exc
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    raise StatementError('Upload the statement as an Excel file (.xlsx) or a CSV file.')


def _heading(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip().lower()


def _heading_row(table):
    """(row index, {column: field}) of the table's heading row. The portal puts account details above it."""
    lookup = {alias: field for field, aliases in STATEMENT_COLUMNS.items() for alias in aliases}
    for index, row in enumerate(table[:40]):
        columns = {}
        for i, heading in enumerate(row):
            field = lookup.get(_heading(heading))
            if field and field not in columns.values():
                columns[i] = field
        if {'receipt', 'time', 'paid_in'} <= set(columns.values()):
            return index, columns
    return None


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, datetime.datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    return str(value).strip()


def read_statement(upload):
    """Money paid in, from an M-Pesa business statement (.xlsx or .csv), as JSON-friendly rows."""
    table = _table(upload)
    found = _heading_row(table)
    if found is None:
        raise StatementError("This doesn't look like an M-Pesa statement: no \"Receipt No.\", \"Completion "
                             'Time" and "Paid In" columns were found. Download the statement from the M-Pesa '
                             'business portal as Excel or CSV and upload it as it is.')
    index, columns = found
    body = table[index + 1:]
    if len(body) > STATEMENT_MAX_ROWS:
        raise StatementError(f'Up to {STATEMENT_MAX_ROWS:,} rows at a time. Download a shorter date range.')
    rows = []
    for line, raw in enumerate(body, start=index + 2):
        values = {field: _cell(raw[i]) if i < len(raw) else '' for i, field in columns.items()}
        receipt = values.get('receipt', '').upper()
        if not receipt or not re.fullmatch(r'[A-Z0-9]{8,12}', receipt):
            continue  # totals, blank lines, footers
        amount = _amount(values.get('paid_in'))
        if not amount or amount <= 0:
            continue  # money paid out, charges and transfers
        status = values.get('status', '').lower()
        if status and status not in ('completed', 'complete', 'success', 'successful'):
            continue
        phone, name = _party(values.get('party', ''), values.get('details', ''))
        rows.append({
            'line': line, 'receipt': receipt, 'time': values.get('time', ''), 'amount': str(amount),
            'account': (values.get('account') or _account_from_details(values.get('details', ''))).strip()[:100],
            'phone': phone, 'name': name, 'details': values.get('details', '')[:200],
        })
    if not rows:
        raise StatementError('No payments received were found in this statement. Check the date range you '
                             'downloaded includes payments from tenants.')
    return rows


def _amount(value):
    try:
        return Decimal(str(value).replace(',', '').replace('KES', '').replace('Ksh', '').strip() or '0')
    except InvalidOperation:
        return None


def _party(party, details):
    """("254712345678", "JOHN DOE") from "254712345678 - JOHN DOE" (numbers are often masked)."""
    text = party or re.sub(r'^.*?\bfrom\s+', '', details, flags=re.I).split(' Acc')[0]
    number, _, name = text.partition(' - ')
    return number.strip()[:30], name.strip()[:100]


def _account_from_details(details):
    match = re.search(r'\bAcc(?:ount)?\.?\s*(?:No\.?)?\s*[:#]?\s*(\S+)', details, re.I)
    return match.group(1) if match else ''


def parse_statement_time(value):
    value = (value or '').strip()
    for fmt in TIME_FORMATS:
        try:
            return timezone.make_aware(datetime.datetime.strptime(value, fmt))
        except ValueError:
            continue
    return None


def _house_key(value):
    return re.sub(r'[^A-Z0-9]', '', (value or '').upper()) or None


def check_statement(rows, account):
    """Each row with what will happen to it: 'match' (a tenant), 'unmatched' (to assign by hand),
    'recorded' (already in Rentipay) or 'error'."""
    from rentals.models import Lease

    from .models import C2BPayment, Payment
    from .mpesa import match_lease

    receipts = [r['receipt'] for r in rows]
    known = set(Payment.all_objects.filter(method=Payment.Method.MPESA, reference__in=receipts)
                .values_list('reference', flat=True))
    known |= set(C2BPayment.objects.filter(trans_id__in=receipts).values_list('trans_id', flat=True))
    # Tenants often type their house number instead of the account number. Live payments don't match on
    # it (nobody checks them first), but every statement row is shown before it's recorded.
    by_house = {}
    for lease in Lease.objects.filter(is_active=True, unit__property__owner=account.owner).select_related('unit'):
        by_house.setdefault(_house_key(lease.unit.label), []).append(lease)
    leases = {}
    seen, checked = set(), []
    for row in rows:
        item = {**row, 'state': '', 'lease': None, 'error': '', 'by_house': False}
        when = parse_statement_time(row['time'])
        if row['receipt'] in known or row['receipt'] in seen:
            item['state'] = 'recorded'
        elif when is None:
            item['state'], item['error'] = 'error', f'Couldn\'t read the date "{row["time"]}".'
        else:
            lease = match_lease(account, row['account'], row['phone'])
            if lease is None and len(by_house.get(_house_key(row['account']), [])) == 1:
                lease, item['by_house'] = by_house[_house_key(row['account'])][0], True
            item['state'] = 'match' if lease else 'unmatched'
            if lease:
                leases[lease.pk] = lease
                item['lease'] = lease
        seen.add(row['receipt'])
        checked.append(item)
    # One query for display details, rather than one per row.
    if leases:
        full = Lease.objects.select_related('tenant', 'unit__property').in_bulk(list(leases))
        for item in checked:
            if item['lease'] is not None:
                item['lease'] = full[item['lease'].pk]
    return checked


def apply_statement(checked, account, actor, send_receipts=False):
    """Record the payments from a checked statement. Returns (matched, unmatched) counts."""
    from audit import utils as audit

    from .models import C2BPayment
    from .mpesa import apply_c2b_payment

    matched = unmatched = 0
    with transaction.atomic():
        for item in checked:
            if item['state'] not in ('match', 'unmatched'):
                continue
            c2b, created = C2BPayment.objects.get_or_create(trans_id=item['receipt'], defaults={
                'source': C2BPayment.Source.STATEMENT, 'account': account,
                'trans_type': 'Statement', 'trans_time': parse_statement_time(item['time']),
                'amount': Decimal(item['amount']), 'shortcode': account.pay_number,
                'bill_ref': item['account'], 'msisdn': item['phone'], 'payer_name': item['name'],
                'raw': {k: item[k] for k in ('line', 'receipt', 'time', 'amount', 'account', 'phone', 'name',
                                              'details')},
            })
            if not created:
                continue  # arrived live while the preview was open
            if item['state'] == 'match':
                c2b.lease = item['lease']
                apply_c2b_payment(c2b, send_receipt=send_receipts)
                c2b.save()
                matched += 1
            else:
                unmatched += 1
        account.statement_imported_at = timezone.now()
        account.save(update_fields=['statement_imported_at', 'updated_at'])
        audit.record(owner=account.owner, actor=actor, action='payment.statement_uploaded',
                     summary=f'Uploaded an M-Pesa statement: {matched} payment{"s" if matched != 1 else ""} '
                             f'recorded, {unmatched} to assign.')
    return matched, unmatched



# --- Cash with staff ------------------------------------------------------------------------------------
# A caretaker who records a cash payment is holding that cash until the landlord confirms receiving it.

def staff_cash(member):
    """What a caretaker has collected, handed over and is holding (from when payments recorded who took them)."""
    from django.db.models import Sum

    from .models import CashHandover, Payment, StaffExpense

    def total(qs):
        return qs.aggregate(t=Sum('amount'))['t'] or Decimal('0')

    collected = total(Payment.objects.filter(method=Payment.Method.CASH, recorded_by=member.user_id,
                                             lease__unit__property__owner=member.owner_id))
    handovers = CashHandover.objects.filter(staff=member)
    confirmed = total(handovers.filter(status=CashHandover.Status.CONFIRMED))
    pending = total(handovers.filter(status=CashHandover.Status.PENDING))
    expenses = StaffExpense.objects.filter(staff=member)
    spent = total(expenses.filter(status=StaffExpense.Status.APPROVED))
    spend_pending = total(expenses.filter(status=StaffExpense.Status.PENDING))
    holding = collected - confirmed - spent
    return {'member': member, 'collected': collected, 'confirmed': confirmed, 'pending': pending,
            'spent': spent, 'spend_pending': spend_pending, 'holding': holding,
            # What they can still hand over or spend: anything already claimed is waiting on the landlord.
            'to_hand_over': max(holding - pending - spend_pending, Decimal('0'))}


def staff_cash_for(owner):
    from accounts.models import StaffMember

    members = StaffMember.objects.filter(owner=owner).select_related('user')
    return [staff_cash(m) for m in members]


class HandoverError(ValueError):
    pass


def record_handover(member, amount, method, reference, note, actor, handed_at=None):
    """A caretaker says they passed cash on; if the landlord records it themselves, it's confirmed at once."""
    from audit import utils as audit

    from .models import CashHandover

    amount = Decimal(amount)
    available = staff_cash(member)['to_hand_over']
    landlord = actor.pk == member.owner_id
    if amount <= 0:
        raise HandoverError('Enter the amount handed over.')
    if not landlord and amount > available:
        raise HandoverError(f'That is more than the cash you are holding (KES {available:,.0f}).')
    handover = CashHandover.objects.create(
        owner_id=member.owner_id, staff=member, amount=amount, method=method, reference=reference[:50],
        note=note[:200], handed_at=handed_at or timezone.now(), recorded_by=actor,
        status=CashHandover.Status.CONFIRMED if landlord else CashHandover.Status.PENDING,
        decided_at=timezone.now() if landlord else None, decided_by=actor if landlord else None)
    who = 'You recorded' if landlord else f'{member.name} (staff) recorded'
    audit.record(owner=member.owner, actor=actor, action='cash.handed_over',
                 summary=f'{who} handing over KES {amount:,.0f} cash ({handover.get_method_display().lower()}'
                         f'{", " + handover.reference if handover.reference else ""}).')
    return handover


def decide_handover(handover, confirm, actor):
    from audit import utils as audit

    from .models import CashHandover

    handover.status = CashHandover.Status.CONFIRMED if confirm else CashHandover.Status.DISPUTED
    handover.decided_at, handover.decided_by = timezone.now(), actor
    handover.save(update_fields=['status', 'decided_at', 'decided_by'])
    audit.record(owner=handover.owner, actor=actor, action='cash.confirmed' if confirm else 'cash.disputed',
                 summary=f'{"Confirmed" if confirm else "Marked as not received"}: KES {handover.amount:,.0f} cash '
                         f'from {handover.staff.name}.')
    return handover


def record_expense(member, amount, description, unit, receipt, spent_on, actor):
    """A caretaker spent some of the cash they hold; it counts once the landlord approves it."""
    from audit import utils as audit
    from listings.utils import prepare_photo

    from .models import StaffExpense

    amount = Decimal(amount)
    available = staff_cash(member)['to_hand_over']
    if amount > available:
        raise HandoverError(f'That is more than the cash you are holding (KES {available:,.0f}).')
    expense = StaffExpense(owner_id=member.owner_id, staff=member, amount=amount, description=description[:200],
                           unit=unit, spent_on=spent_on, recorded_by=actor)
    if receipt:
        content = prepare_photo(receipt)  # resized, location data removed
        expense.receipt.save(content.name, content, save=False)
    expense.save()
    audit.record(owner=member.owner, actor=actor, action='cash.spent',
                 summary=f'{member.name} (staff) recorded spending KES {amount:,.0f} of cash on "{description}"'
                         f'{" (" + unit.label + ")" if unit else ""}. Waiting for approval.')
    return expense


def decide_expense(expense, approve, actor, reason=''):
    from audit import utils as audit

    from .models import StaffExpense

    expense.status = StaffExpense.Status.APPROVED if approve else StaffExpense.Status.REJECTED
    expense.reject_reason = '' if approve else reason[:200]
    expense.decided_at, expense.decided_by = timezone.now(), actor
    expense.save(update_fields=['status', 'reject_reason', 'decided_at', 'decided_by'])
    if approve:
        expense_from_staff_spending(expense, actor)
    audit.record(owner=expense.owner, actor=actor, action='cash.spending_approved' if approve else 'cash.spending_rejected',
                 summary=f'{"Approved" if approve else "Did not approve"} KES {expense.amount:,.0f} spent by '
                         f'{expense.staff.name} on "{expense.description}"{"" if approve else ": " + reason}.')
    return expense



# --- Expenses and profit ------------------------------------------------------------------------------

def _same_day_next_month(day, anchor_day):
    import calendar

    year, month = (day.year + 1, 1) if day.month == 12 else (day.year, day.month + 1)
    return day.replace(year=year, month=month, day=min(anchor_day, calendar.monthrange(year, month)[1]))


def create_recurring_expenses(today=None):
    """Add this month's copy of every expense set to repeat monthly (and catch up missed months). Returns count."""
    from .models import Expense

    today = today or timezone.localdate()
    made = 0
    for original in Expense.objects.filter(repeats_monthly=True, repeat_of__isnull=True):
        # Voided copies count too, so a month the landlord cancelled isn't added again.
        copies = Expense.all_objects.filter(repeat_of=original).values_list('paid_on', flat=True)
        latest = max([original.paid_on, *copies])
        due = _same_day_next_month(latest, original.paid_on.day)
        while due <= today:
            Expense.objects.create(owner_id=original.owner_id, property_id=original.property_id,
                                   unit_id=original.unit_id, category=original.category,
                                   description=original.description, amount=original.amount, paid_on=due,
                                   method=original.method, repeat_of=original)
            made += 1
            due = _same_day_next_month(due, original.paid_on.day)
    return made


def expense_from_staff_spending(staff_expense, actor):
    """Approved caretaker spending is an expense too (paid from the rent cash they held)."""
    from .models import Expense

    unit = staff_expense.unit
    expense, _ = Expense.objects.get_or_create(staff_expense=staff_expense, defaults={
        'owner_id': staff_expense.owner_id, 'unit': unit, 'property_id': unit.property_id if unit else None,
        'category': Expense.Category.REPAIRS, 'description': f'{staff_expense.description} ({staff_expense.staff.name})',
        'amount': staff_expense.amount, 'paid_on': staff_expense.spent_on, 'method': Expense.Method.STAFF_CASH,
        'recorded_by': actor,
    })
    return expense


def expenses_by_property(owner, start, end):
    """{property id (None = shared costs): total spent} between start (inclusive) and end (exclusive)."""
    from django.db.models import Sum

    from .models import Expense

    return dict(Expense.objects.filter(owner=owner, paid_on__gte=start, paid_on__lt=end)
                .values_list('property').annotate(total=Sum('amount')))
