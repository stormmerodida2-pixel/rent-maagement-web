"""Importing units and tenants from a spreadsheet."""

import datetime
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from reports.utils import month_summary

from ..models import Lease, Property, RentCharge, Tenant, Unit
from . import make_lease

HEAD = ['Unit', 'Type', 'Bedrooms', 'Rent', 'Tenant name', 'Phone', 'Email', 'Rent day', 'Moved in', 'Deposit',
        'Balance owed now', 'Language']
TODAY = datetime.date.today()


def xlsx(rows, head=HEAD):
    book = Workbook()
    sheet = book.active
    sheet.append(head)
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return SimpleUploadedFile('units.xlsx', out.getvalue())


class ImportTests(TestCase):
    def setUp(self):
        self.owner = get_user_model().objects.create_user('landlord', password='pw')
        self.prop = Property.objects.create(owner=self.owner, name='Riverside')
        self.client.force_login(self.owner)
        self.url = reverse('rentals:property_import')

    def preview(self, upload, **extra):
        return self.client.post(self.url, {'action': 'preview', 'property': self.prop.pk, 'file': upload, **extra})

    def confirm(self):
        return self.client.post(self.url, {'action': 'confirm'})

    def test_template_download(self):
        resp = self.client.get(reverse('rentals:import_template'))
        self.assertEqual(resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        sheet = load_workbook(io.BytesIO(resp.content)).active
        # Older sheets with a Bedrooms column (HEAD, used below) are still read.
        self.assertEqual([c.value for c in sheet[1]], [h for h in HEAD if h != 'Bedrooms'])
        self.assertEqual(sheet.data_validations.dataValidation[0].formula1, '"Single room,Double room,One bedroom,Two bedroom,Three bedroom,Shop"')

    def test_preview_checks_rows_and_saves_nothing(self):
        occupied = make_lease(self.owner, label='Z1')
        Unit.objects.filter(pk=occupied.unit_id).update(property=self.prop)
        resp = self.preview(xlsx([
            ['A1', '1-bedroom', 1, 18000, 'Wanjiku Kamau', '0712 345 678', '', 1, '01/06/2024', 18000, 3000, 'Kiswahili'],
            ['A2', 'Bedsitter', 0, 'KES 12,000', '', '', '', '', '', '', '', ''],
            ['A3', '', 1, 15000, 'Bad Phone', '12345', '', '', '', '', '', ''],
            ['A1', '', 1, 15000, '', '', '', '', '', '', '', ''],
            ['Z1', '', 1, 15000, 'Someone Else', '0722000111', '', '', '', '', '', ''],
            ['A4', '', 1, '', '', '', '', '', '', '', '', ''],
        ]))
        self.assertContains(resp, '2 rows ready')
        self.assertContains(resp, 'Wanjiku Kamau moves in')
        self.assertContains(resp, 'owes KES 3,000')
        self.assertContains(resp, 'isn&#x27;t a Kenyan mobile number')
        self.assertContains(resp, 'Unit A1 appears twice')
        self.assertContains(resp, 'Z1 already has a tenant')
        self.assertContains(resp, 'Rent is missing')
        self.assertFalse(Unit.objects.filter(label='A1').exists())  # nothing saved yet

    def test_confirm_carries_over_balances_without_double_billing(self):
        self.preview(xlsx([
            ['A1', '1-bedroom', 1, 18000, 'Wanjiku Kamau', '0712 345 678', 'w@example.com', 1, '01/06/2024', 18000,
             3000, 'Kiswahili'],
            ['A2', 'Bedsitter', 0, 12000, '', '', '', '', '', '', '', ''],
        ]))
        resp = self.confirm()
        self.assertRedirects(resp, reverse('rentals:property_detail', args=[self.prop.pk]))
        a2 = Unit.objects.get(property=self.prop, label='A2')
        self.assertEqual((a2.default_rent, a2.unit_type, a2.active_lease()), (Decimal('12000'), 'single', None))
        self.assertEqual(Unit.objects.get(property=self.prop, label='A1').unit_type, 'one_bedroom')
        lease = Lease.objects.get(unit__label='A1')
        self.assertEqual((lease.monthly_rent, lease.due_day, lease.start_date), (Decimal('18000'), 1, datetime.date(2024, 6, 1)))
        self.assertEqual(lease.tenant.language, 'sw')
        self.assertEqual(lease.balance, Decimal('3000'))  # exactly what they owed; the deposit isn't added
        self.assertEqual(lease.deposit_held, Decimal('18000'))
        # Rent day 1 has passed, so this month's rent was part of the balance: billing it again does nothing.
        lease.ensure_charge_for(TODAY)
        self.assertEqual(lease.balance, Decimal('3000'))
        self.assertTrue(RentCharge.all_objects.filter(lease=lease, kind='rent', voided_at__isnull=False).exists())
        # The carried-over balance isn't counted as this month's billing in reports.
        self.assertEqual(month_summary(self.owner, TODAY.replace(day=1))['billed'], Decimal('0'))
        statement = self.client.get(reverse('rentals:lease_detail', args=[lease.pk]))
        self.assertContains(statement, 'Balance brought forward')
        self.assertContains(statement, 'Deposit paid before Rentipay')
        self.assertIsNone(self.client.session.get('import_preview'))

    def test_csv_and_a_new_property(self):
        csv = SimpleUploadedFile('units.csv', b'Unit,Rent,Tenant name,Phone\nB1,"20,000",Otieno Mwangi,0733000222\n')
        self.client.post(self.url, {'action': 'preview', 'property': 'new', 'new_property': 'Sunrise', 'file': csv})
        self.confirm()
        prop = Property.objects.get(owner=self.owner, name='Sunrise')
        self.assertEqual(Lease.objects.get(unit__property=prop).tenant.full_name, 'Otieno Mwangi')

    def test_existing_contact_is_reused(self):
        Tenant.objects.create(owner=self.owner, full_name='Jane W', phone='0712345678')
        self.preview(xlsx([['A1', '', 1, 18000, 'Jane Wanjiru', '0712345678', '', '', '', '', '', '']]))
        self.confirm()
        self.assertEqual(Tenant.objects.filter(owner=self.owner).count(), 1)

    def test_problems_with_the_file_itself(self):
        resp = self.preview(xlsx([['A1', 18000]], head=['Room number', 'Amount']))
        self.assertContains(resp, 'including &quot;Unit&quot; and &quot;Rent&quot;')
        resp = self.preview(SimpleUploadedFile('units.pdf', b'%PDF'))
        self.assertContains(resp, 'Upload an Excel file')

    def test_only_into_your_own_properties(self):
        other = Property.objects.create(owner=get_user_model().objects.create_user('other'), name='Theirs')
        resp = self.client.post(self.url, {'action': 'preview', 'property': other.pk,
                                           'file': xlsx([['A1', '', 1, 18000, '', '', '', '', '', '', '', '']])})
        self.assertContains(resp, 'Choose a property')
        self.assertFalse(Unit.objects.filter(property=other).exists())
